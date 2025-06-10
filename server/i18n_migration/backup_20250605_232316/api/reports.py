"""Reports API endpoints."""

from flask import Blueprint, request, jsonify, send_file
from flask_jwt_extended import jwt_required, get_jwt_identity
from datetime import datetime, timedelta
from sqlalchemy import func
import os
import io
import csv
import json
from openpyxl import Workbook

from app.extensions import db
from app.models.user import User
from app.models.report import Report, ReportSchedule
from app.models.tenant import Tenant

from app.utils.logging import logger

reports_bp = Blueprint('reports', __name__)

# Demo data için değişkenler
DEMO_REPORT_TYPES = ['beneficiary', 'program', 'trainer', 'analytics', 'performance']
DEMO_REPORT_NAMES = {
    'beneficiary': ['Beneficiary Progress Report', 'Monthly Beneficiary Summary', 'Beneficiary Test Results'],
    'program': ['Program Performance Report', 'Program Enrollment Summary', 'Program Completion Analysis'],
    'trainer': ['Trainer Activity Report', 'Trainer Performance Summary', 'Trainer Load Analysis'],
    'analytics': ['Overall Analytics Report', 'KPI Dashboard Summary', 'Trend Analysis Report'],
    'performance': ['Performance Metrics Report', 'Test Score Analysis', 'Learning Outcomes Report']
}
DEMO_REPORT_DESCRIPTIONS = {
    'beneficiary': 'Comprehensive report on beneficiary progress and development',
    'program': 'Detailed analysis of program performance and enrollment metrics',
    'trainer': 'Summary of trainer activities and performance indicators',
    'analytics': 'Overall system analytics and key performance indicators',
    'performance': 'Detailed performance metrics and test score analysis'
}

@reports_bp.route('/reports/recent', methods=['GET'])
@jwt_required()
def get_recent_reports():
    """Get recent reports for the current user."""
    user_id = get_jwt_identity()
    user = User.query.get(user_id)
    
    # Query recent reports
    query = Report.query.filter_by(created_by_id=user_id)
    
    # For admins, show all tenant reports
    if user.role in ['super_admin', 'tenant_admin']:
        query = Report.query.filter_by(tenant_id=user.tenant_id)
    
    # Get last 10 reports
    recent_reports = query.order_by(Report.created_at.desc()).limit(10).all()
    
    return jsonify([report.to_dict() for report in recent_reports]), 200

@reports_bp.route('/reports/saved', methods=['GET'])
@jwt_required()
def get_saved_reports():
    """Get saved reports for the current user."""
    user_id = get_jwt_identity()
    user = User.query.get(user_id)
    
    # Query saved reports
    query = Report.query.filter_by(status='completed')
    
    if user.role in ['super_admin', 'tenant_admin']:
        query = query.filter_by(tenant_id=user.tenant_id)
    else:
        query = query.filter_by(created_by_id=user_id)
    
    saved_reports = query.order_by(Report.updated_at.desc()).all()
    
    return jsonify([report.to_dict() for report in saved_reports]), 200

@reports_bp.route('/reports/scheduled', methods=['GET'])
@jwt_required()
def get_scheduled_reports():
    """Get scheduled reports for the current user."""
    user_id = get_jwt_identity()
    user = User.query.get(user_id)
    
    # Query scheduled reports
    query = ReportSchedule.query.join(Report).filter(ReportSchedule.is_active == True)
    
    if user.role in ['super_admin', 'tenant_admin']:
        query = query.filter(Report.tenant_id == user.tenant_id)
    else:
        query = query.filter(Report.created_by_id == user_id)
    
    scheduled_reports = query.all()
    
    return jsonify([schedule.to_dict() for schedule in scheduled_reports]), 200


@reports_bp.route('/reports', methods=['POST'])
@jwt_required()
def create_report():
    """Create a new report."""
    user_id = get_jwt_identity()
    user = User.query.get(user_id)
    data = request.get_json()
    
    try:
        report = Report(
            name=data['name'],
            description=data.get('description'),
            type=data['type'],
            format=data.get('format', 'pdf'),
            parameters=data.get('parameters', {}),
            created_by_id=user_id,
            tenant_id=user.tenant_id
        )
        
        db.session.add(report)
        db.session.commit()
        
        return jsonify(report.to_dict()), 201
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500


@reports_bp.route('/reports/<int:report_id>', methods=['GET'])
@jwt_required()
def get_report(report_id):
    """Get a specific report."""
    user_id = get_jwt_identity()
    user = User.query.get(user_id)
    
    report = Report.query.get_or_404(report_id)
    
    # Check permissions
    if user.role not in ['super_admin', 'tenant_admin']:
        if report.created_by_id != user_id:
            return jsonify({'error': 'Unauthorized'}), 403
    
    return jsonify(report.to_dict()), 200


@reports_bp.route('/reports/<int:report_id>', methods=['PUT'])
@jwt_required()
def update_report(report_id):
    """Update a report."""
    user_id = get_jwt_identity()
    user = User.query.get(user_id)
    data = request.get_json()
    
    report = Report.query.get_or_404(report_id)
    
    # Check permissions
    if user.role not in ['super_admin', 'tenant_admin']:
        if report.created_by_id != user_id:
            return jsonify({'error': 'Unauthorized'}), 403
    
    # Update report
    report.name = data.get('name', report.name)
    report.description = data.get('description', report.description)
    report.parameters = data.get('parameters', report.parameters)
    
    db.session.commit()
    
    return jsonify(report.to_dict()), 200


@reports_bp.route('/reports/<int:report_id>/run', methods=['POST'])
@jwt_required()
def run_report(report_id):
    """Run a report and generate output."""
    user_id = get_jwt_identity()
    user = User.query.get(user_id)
    
    report = Report.query.get_or_404(report_id)
    
    # Check permissions
    if user.role not in ['super_admin', 'tenant_admin']:
        if report.created_by_id != user_id:
            return jsonify({'error': 'Unauthorized'}), 403
    
    try:
        # Update report status
        report.status = 'generating'
        db.session.commit()
        
        # Generate report based on type
        if report.type == 'beneficiary':
            data = generate_beneficiary_report(report.parameters)
        elif report.type == 'trainer':
            data = generate_trainer_report(report.parameters)
        elif report.type == 'program':
            data = generate_program_report(report.parameters)
        elif report.type == 'performance':
            data = generate_performance_report(report.parameters)
        else:
            data = generate_general_report(report.parameters)
        
        # Create file based on format
        if report.format == 'pdf':
            file_path = generate_pdf_report(data, report)
        elif report.format == 'xlsx':
            file_path = generate_excel_report(data, report)
        else:  # csv
            file_path = generate_csv_report(data, report)
        
        # Update report
        report.status = 'completed'
        report.file_path = file_path
        report.file_size = os.path.getsize(file_path)
        report.last_generated = datetime.utcnow()
        report.run_count += 1
        
        db.session.commit()
        
        return jsonify(report.to_dict()), 200
        
    except Exception as e:
        report.status = 'failed'
        db.session.commit()
        return jsonify({'error': str(e)}), 500


@reports_bp.route('/reports/<int:report_id>/download', methods=['GET'])
@jwt_required()
def download_report(report_id):
    """Download a generated report."""
    user_id = get_jwt_identity()
    user = User.query.get(user_id)
    
    report = Report.query.get_or_404(report_id)
    
    # Check permissions
    if user.role not in ['super_admin', 'tenant_admin']:
        if report.created_by_id != user_id:
            return jsonify({'error': 'Unauthorized'}), 403
    
    if not report.file_path or not os.path.exists(report.file_path):
        return jsonify({'error': 'Report file not found'}), 404
    
    return send_file(
        report.file_path,
        as_attachment=True,
        download_name=f"{report.name}_{datetime.now().strftime('%Y%m%d')}.{report.format}"
    )


@reports_bp.route('/reports/<int:report_id>/download-pdf', methods=['GET'])
@jwt_required()
def download_report_pdf(report_id):
    """Download a report as PDF."""
    from app.utils.pdf_generator import generate_report_pdf
    
    user_id = get_jwt_identity()
    user = User.query.get(user_id)
    
    report = Report.query.get_or_404(report_id)
    
    # Check permissions
    if user.role not in ['super_admin', 'tenant_admin']:
        if report.created_by_id != user_id:
            return jsonify({'error': 'Unauthorized'}), 403
    
    try:
        # Generate report data based on type
        if report.type == 'beneficiary':
            data = generate_beneficiary_report(report.parameters)
        elif report.type == 'trainer':
            data = generate_trainer_report(report.parameters)
        elif report.type == 'program':
            data = generate_program_report(report.parameters)
        elif report.type == 'performance':
            data = generate_performance_report(report.parameters)
        else:
            data = generate_general_report(report.parameters)
        
        # Prepare template data
        template_data = {
            'title': report.name,
            'description': report.description,
            'generated_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'generated_by': f"{user.first_name} {user.last_name}",
            'data': data,
            'report_type': report.type
        }
        
        # Generate PDF
        pdf_data = generate_report_pdf(template_data)
        
        # Return PDF as download
        return send_file(
            io.BytesIO(pdf_data),
            mimetype='application/pdf',
            as_attachment=True,
            download_name=f"{report.name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        )
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500


def generate_beneficiary_report(parameters):
    """Generate beneficiary report data."""
    from app.models.beneficiary import Beneficiary
    from app.models.test import TestSession
    
    query = Beneficiary.query
    
    # Apply filters from parameters
    if 'beneficiary_ids' in parameters:
        query = query.filter(Beneficiary.id.in_(parameters['beneficiary_ids']))
    
    beneficiaries = query.all()
    
    data = []
    for beneficiary in beneficiaries:
        # Get test scores
        test_scores = db.session.query(func.avg(TestSession.score)).filter(
            TestSession.beneficiary_id == beneficiary.id,
            TestSession.status == 'completed'
        ).scalar() or 0
        
        data.append({
            'Name': f"{beneficiary.first_name} {beneficiary.last_name}",
            'Email': beneficiary.email,
            'Status': beneficiary.status,
            'Average Score': round(test_scores, 2),
            'Created': beneficiary.created_at.strftime('%Y-%m-%d') if beneficiary.created_at else ''
        })
    
    return data


def generate_trainer_report(parameters):
    """Generate trainer report data."""
    from app.models.beneficiary import Beneficiary
    
    query = User.query.filter_by(role='trainer')
    
    # Apply filters from parameters
    if 'trainer_ids' in parameters:
        query = query.filter(User.id.in_(parameters['trainer_ids']))
    
    trainers = query.all()
    
    data = []
    for trainer in trainers:
        beneficiary_count = Beneficiary.query.filter_by(trainer_id=trainer.id).count()
        
        data.append({
            'Name': f"{trainer.first_name} {trainer.last_name}",
            'Email': trainer.email,
            'Beneficiaries': beneficiary_count,
            'Active': 'Yes' if trainer.is_active else 'No',
            'Last Login': trainer.last_login.strftime('%Y-%m-%d %H:%M') if trainer.last_login else ''
        })
    
    return data


def generate_program_report(parameters):
    """Generate program report data."""
    from app.models.program import Program, ProgramEnrollment, Module, Session, SessionAttendance
    
    query = Program.query
    
    # Apply filters
    if 'program_ids' in parameters:
        query = query.filter(Program.id.in_(parameters['program_ids']))
    
    if 'status' in parameters:
        query = query.filter_by(status=parameters['status'])
    
    programs = query.all()
    
    data = []
    for program in programs:
        enrollment_count = ProgramEnrollment.query.filter_by(program_id=program.id).count()
        active_enrollments = ProgramEnrollment.query.filter_by(
            program_id=program.id,
            status='enrolled'
        ).count()
        completed_enrollments = ProgramEnrollment.query.filter_by(
            program_id=program.id,
            status='completed'
        ).count()
        
        # Calculate average attendance
        sessions = Session.query.filter_by(program_id=program.id).all()
        total_attendance = 0
        total_possible = 0
        
        for session in sessions:
            attendances = SessionAttendance.query.filter_by(session_id=session.id).all()
            present_count = sum(1 for a in attendances if a.attendance_status == 'present')
            total_attendance += present_count
            total_possible += len(attendances)
        
        attendance_rate = (total_attendance / total_possible * 100) if total_possible > 0 else 0
        
        data.append({
            'Program Name': program.name,
            'Code': program.code,
            'Status': program.status,
            'Total Enrollments': enrollment_count,
            'Active': active_enrollments,
            'Completed': completed_enrollments,
            'Completion Rate': f"{(completed_enrollments / enrollment_count * 100):.1f}%" if enrollment_count > 0 else "0%",
            'Attendance Rate': f"{attendance_rate:.1f}%",
            'Start Date': program.start_date.strftime('%Y-%m-%d') if program.start_date else '',
            'End Date': program.end_date.strftime('%Y-%m-%d') if program.end_date else ''
        })
    
    return data


def generate_performance_report(parameters):
    """Generate performance report data."""
    from app.models.test import TestSession
    from app.models.beneficiary import Beneficiary
    
    # Get date range
    start_date = parameters.get('start_date', datetime.now() - timedelta(days=30))
    end_date = parameters.get('end_date', datetime.now())
    
    # Query test sessions
    sessions = TestSession.query.filter(
        TestSession.completed_at.between(start_date, end_date),
        TestSession.status == 'completed'
    ).join(Beneficiary).all()
    
    data = []
    for session in sessions:
        data.append({
            'Date': session.completed_at.strftime('%Y-%m-%d %H:%M'),
            'Beneficiary': f"{session.beneficiary.first_name} {session.beneficiary.last_name}",
            'Test': session.test.title if session.test else 'Unknown',
            'Score': session.score,
            'Duration': session.duration,
            'Status': session.status
        })
    
    return data


def generate_general_report(parameters):
    """Generate general report data."""
    # Default implementation
    return []


def generate_excel_report(data, report):
    """Generate Excel report file."""
    wb = Workbook()
    ws = wb.active
    ws.title = report.name[:31]  # Excel sheet name limit
    
    if data:
        # Add headers
        headers = list(data[0].keys())
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
        
        # Add data
        for row_num, row_data in enumerate(data, 2):
            for col, header in enumerate(headers, 1):
                ws.cell(row=row_num, column=col, value=row_data.get(header, ''))
    
    # Create directory if not exists
    reports_dir = os.path.join('app', 'static', 'reports')
    os.makedirs(reports_dir, exist_ok=True)
    
    # Save file
    filename = f"report_{report.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    file_path = os.path.join(reports_dir, filename)
    wb.save(file_path)
    
    return file_path


def generate_csv_report(data, report):
    """Generate CSV report file."""
    # Create directory if not exists
    reports_dir = os.path.join('app', 'static', 'reports')
    os.makedirs(reports_dir, exist_ok=True)
    
    # Save file
    filename = f"report_{report.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    file_path = os.path.join(reports_dir, filename)
    
    if data:
        headers = list(data[0].keys())
        with open(file_path, 'w', newline='', encoding='utf-8') as csvfile:
            writer = csv.DictWriter(csvfile, fieldnames=headers)
            writer.writeheader()
            writer.writerows(data)
    else:
        # Create empty file
        open(file_path, 'a').close()
    
    return file_path


def generate_pdf_report(data, report):
    """Generate PDF report file."""
    from app.utils.pdf_generator import generate_report_pdf
    
    # Create directory if not exists
    reports_dir = os.path.join('app', 'static', 'reports')
    os.makedirs(reports_dir, exist_ok=True)
    
    # Save file
    filename = f"report_{report.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    file_path = os.path.join(reports_dir, filename)
    
    # Prepare template data
    template_data = {
        'title': report.name,
        'description': report.description,
        'generated_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'generated_by': f"{report.created_by.first_name} {report.created_by.last_name}",
        'data': data,
        'report_type': report.type
    }
    
    # Generate PDF
    pdf_data = generate_report_pdf(template_data)
    
    # Save PDF to file
    with open(file_path, 'wb') as pdf_file:
        pdf_file.write(pdf_data)
    
    return file_path


@reports_bp.route('/reports/demo', methods=['POST'])
@jwt_required()
def create_demo_reports():
    """Create demo reports for testing."""
    user_id = get_jwt_identity()
    user = User.query.get(user_id)
    
    try:
        created_reports = []
        
        # Create 3 reports for each type
        for report_type in DEMO_REPORT_TYPES:
            for i, name in enumerate(DEMO_REPORT_NAMES[report_type]):
                report = Report(
                    name=name,
                    description=DEMO_REPORT_DESCRIPTIONS[report_type],
                    type=report_type,
                    format=['pdf', 'xlsx', 'csv'][i % 3],
                    status='completed',
                    parameters={
                        'date_range': 'last_30_days',
                        'filters': {'active': True}
                    },
                    created_by_id=user_id,
                    tenant_id=user.tenant_id,
                    run_count=i + 1,
                    last_generated=datetime.now() - timedelta(days=i)
                )
                
                # Set some as draft or generating
                if i % 5 == 0:
                    report.status = 'draft'
                elif i % 7 == 0:
                    report.status = 'generating'
                
                db.session.add(report)
                created_reports.append(report)
        
        # Create scheduled reports
        for i, report in enumerate(created_reports[:5]):
            schedule = ReportSchedule(
                report_id=report.id,
                frequency=['daily', 'weekly', 'monthly'][i % 3],
                schedule_time=datetime.now().time(),
                recipients=['admin@bdc.com', f'user{i}@bdc.com'],
                recipients_count=2,
                is_active=i % 2 == 0,
                status='active' if i % 2 == 0 else 'paused',
                next_run=datetime.now() + timedelta(days=i),
                last_run=datetime.now() - timedelta(days=i) if i > 0 else None
            )
            db.session.add(schedule)
        
        db.session.commit()
        
        return jsonify({
            'message': 'Demo reports created successfully',
            'count': len(created_reports)
        }), 201
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500


@reports_bp.route('/reports/<int:report_id>/export', methods=['GET'])
@jwt_required()
def export_report(report_id):
    """Export report in different formats."""
    user_id = get_jwt_identity()
    user = User.query.get(user_id)
    format = request.args.get('format', 'pdf')
    
    report = Report.query.get_or_404(report_id)
    
    # Check permissions
    if user.role not in ['super_admin', 'tenant_admin']:
        if report.created_by_id != user_id:
            return jsonify({'error': 'Unauthorized'}), 403
    
    try:
        # Generate report data based on type
        if report.type == 'beneficiary':
            data = generate_beneficiary_report(report.parameters)
        elif report.type == 'trainer':
            data = generate_trainer_report(report.parameters)
        elif report.type == 'program':
            data = generate_program_report(report.parameters)
        elif report.type == 'performance':
            data = generate_performance_report(report.parameters)
        else:
            data = generate_general_report(report.parameters)
        
        # Export in requested format
        if format == 'pdf':
            file_path = generate_pdf_report(data, report)
        elif format == 'xlsx':
            file_path = generate_excel_report(data, report)
        else:  # csv
            file_path = generate_csv_report(data, report)
        
        return send_file(
            file_path,
            as_attachment=True,
            download_name=f"{report.name}_{datetime.now().strftime('%Y%m%d')}.{format}"
        )
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@reports_bp.route('/reports/fields/<report_type>', methods=['GET'])
@jwt_required()
def get_report_fields(report_type):
    """Get available fields for a report type."""
    
    # Define fields for each report type
    fields_by_type = {
        'beneficiary': [
            {'id': 'name', 'name': 'Full Name', 'description': 'Beneficiary full name'},
            {'id': 'email', 'name': 'Email', 'description': 'Beneficiary email address'},
            {'id': 'status', 'name': 'Status', 'description': 'Current beneficiary status'},
            {'id': 'trainer', 'name': 'Assigned Trainer', 'description': 'Trainer assigned to beneficiary'},
            {'id': 'created_date', 'name': 'Created Date', 'description': 'Date beneficiary was added'},
            {'id': 'test_score', 'name': 'Average Test Score', 'description': 'Average score across all tests'},
            {'id': 'progress', 'name': 'Progress', 'description': 'Overall progress percentage'},
            {'id': 'department', 'name': 'Department', 'description': 'Department classification'},
            {'id': 'notes', 'name': 'Notes', 'description': 'Additional notes and comments'}
        ],
        'program': [
            {'id': 'name', 'name': 'Program Name', 'description': 'Name of the training program'},
            {'id': 'code', 'name': 'Program Code', 'description': 'Unique program identifier'},
            {'id': 'status', 'name': 'Status', 'description': 'Current program status'},
            {'id': 'start_date', 'name': 'Start Date', 'description': 'Program start date'},
            {'id': 'end_date', 'name': 'End Date', 'description': 'Program end date'},
            {'id': 'enrollment_count', 'name': 'Enrollment Count', 'description': 'Number of enrolled beneficiaries'},
            {'id': 'completion_rate', 'name': 'Completion Rate', 'description': 'Program completion percentage'},
            {'id': 'attendance_rate', 'name': 'Attendance Rate', 'description': 'Average attendance rate'},
            {'id': 'description', 'name': 'Description', 'description': 'Program description'}
        ],
        'trainer': [
            {'id': 'name', 'name': 'Trainer Name', 'description': 'Full name of the trainer'},
            {'id': 'email', 'name': 'Email', 'description': 'Trainer email address'},
            {'id': 'beneficiary_count', 'name': 'Beneficiary Count', 'description': 'Number of assigned beneficiaries'},
            {'id': 'active_status', 'name': 'Active Status', 'description': 'Whether trainer is active'},
            {'id': 'last_login', 'name': 'Last Login', 'description': 'Last login timestamp'},
            {'id': 'specialization', 'name': 'Specialization', 'description': 'Areas of expertise'},
            {'id': 'programs', 'name': 'Assigned Programs', 'description': 'Programs trainer is involved in'},
            {'id': 'performance_rating', 'name': 'Performance Rating', 'description': 'Average performance rating'}
        ],
        'analytics': [
            {'id': 'metric_name', 'name': 'Metric Name', 'description': 'Name of the metric'},
            {'id': 'value', 'name': 'Value', 'description': 'Metric value'},
            {'id': 'change', 'name': 'Change', 'description': 'Change from previous period'},
            {'id': 'trend', 'name': 'Trend', 'description': 'Trend direction'},
            {'id': 'date', 'name': 'Date', 'description': 'Date of the metric'},
            {'id': 'category', 'name': 'Category', 'description': 'Metric category'}
        ],
        'performance': [
            {'id': 'beneficiary_name', 'name': 'Beneficiary Name', 'description': 'Full name of the beneficiary'},
            {'id': 'test_name', 'name': 'Test Name', 'description': 'Name of the test'},
            {'id': 'score', 'name': 'Score', 'description': 'Test score'},
            {'id': 'date', 'name': 'Date', 'description': 'Test completion date'},
            {'id': 'duration', 'name': 'Duration', 'description': 'Time taken to complete'},
            {'id': 'status', 'name': 'Status', 'description': 'Test completion status'},
            {'id': 'feedback', 'name': 'Feedback', 'description': 'Test feedback and comments'},
            {'id': 'improvement', 'name': 'Improvement', 'description': 'Improvement from previous test'}
        ]
    }
    
    fields = fields_by_type.get(report_type, [])
    
    return jsonify(fields), 200


@reports_bp.route('/reports/filters/<report_type>', methods=['GET'])
@jwt_required()
def get_report_filters(report_type):
    """Get available filters for a report type."""
    
    # Define filters for each report type
    filters_by_type = {
        'beneficiary': [
            {
                'id': 'status',
                'name': 'Status',
                'type': 'select',
                'options': [
                    {'value': 'active', 'label': 'Active'},
                    {'value': 'inactive', 'label': 'Inactive'},
                    {'value': 'completed', 'label': 'Completed'},
                    {'value': 'pending', 'label': 'Pending'}
                ]
            },
            {
                'id': 'trainer',
                'name': 'Trainer',
                'type': 'multiselect',
                'options': []  # Would be populated dynamically
            },
            {
                'id': 'test_score_range',
                'name': 'Test Score Range',
                'type': 'range',
                'options': {'min': 0, 'max': 100}
            },
            {
                'id': 'created_date',
                'name': 'Created Date',
                'type': 'date',
                'options': {}
            }
        ],
        'program': [
            {
                'id': 'status',
                'name': 'Status',
                'type': 'select',
                'options': [
                    {'value': 'active', 'label': 'Active'},
                    {'value': 'completed', 'label': 'Completed'},
                    {'value': 'upcoming', 'label': 'Upcoming'},
                    {'value': 'cancelled', 'label': 'Cancelled'}
                ]
            },
            {
                'id': 'enrollment_range',
                'name': 'Enrollment Count',
                'type': 'range',
                'options': {'min': 0, 'max': 100}
            },
            {
                'id': 'date_range',
                'name': 'Date Range',
                'type': 'date',
                'options': {}
            }
        ],
        'trainer': [
            {
                'id': 'active_status',
                'name': 'Active Status',
                'type': 'select',
                'options': [
                    {'value': 'active', 'label': 'Active'},
                    {'value': 'inactive', 'label': 'Inactive'}
                ]
            },
            {
                'id': 'beneficiary_count',
                'name': 'Beneficiary Count',
                'type': 'range',
                'options': {'min': 0, 'max': 50}
            },
            {
                'id': 'programs',
                'name': 'Programs',
                'type': 'multiselect',
                'options': []  # Would be populated dynamically
            }
        ],
        'analytics': [
            {
                'id': 'metric_category',
                'name': 'Metric Category',
                'type': 'multiselect',
                'options': [
                    {'value': 'engagement', 'label': 'Engagement'},
                    {'value': 'performance', 'label': 'Performance'},
                    {'value': 'completion', 'label': 'Completion'},
                    {'value': 'satisfaction', 'label': 'Satisfaction'}
                ]
            },
            {
                'id': 'date_range',
                'name': 'Date Range',
                'type': 'date',
                'options': {}
            }
        ],
        'performance': [
            {
                'id': 'test_name',
                'name': 'Test Name',
                'type': 'multiselect',
                'options': []  # Would be populated dynamically
            },
            {
                'id': 'score_range',
                'name': 'Score Range',
                'type': 'range',
                'options': {'min': 0, 'max': 100}
            },
            {
                'id': 'status',
                'name': 'Status',
                'type': 'select',
                'options': [
                    {'value': 'completed', 'label': 'Completed'},
                    {'value': 'in_progress', 'label': 'In Progress'},
                    {'value': 'abandoned', 'label': 'Abandoned'}
                ]
            },
            {
                'id': 'date_range',
                'name': 'Date Range',
                'type': 'date',
                'options': {}
            }
        ]
    }
    
    filters = filters_by_type.get(report_type, [])
    
    # Populate dynamic options
    if report_type == 'beneficiary' and filters:
        # Get trainers for filter
        trainers = User.query.filter_by(role='trainer').all()
        for filter in filters:
            if filter['id'] == 'trainer':
                filter['options'] = [
                    {'value': str(t.id), 'label': f"{t.first_name} {t.last_name}"} 
                    for t in trainers
                ]
    
    return jsonify(filters), 200


@reports_bp.route('/reports/preview', methods=['POST'])
@jwt_required()
def preview_report():
    """Generate a preview of the report."""
    data = request.get_json()
    
    try:
        # Generate preview data based on report type
        preview_data = {
            'sections': []
        }
        
        if data['type'] == 'beneficiary':
            # Summary section
            preview_data['sections'].append({
                'title': 'Summary',
                'type': 'summary',
                'metrics': [
                    {'id': 'total', 'name': 'Total Beneficiaries', 'value': '125', 'change': 12},
                    {'id': 'active', 'name': 'Active', 'value': '98', 'change': 5},
                    {'id': 'avg_score', 'name': 'Average Score', 'value': '82.5%', 'change': 3}
                ]
            })
            
            # Table section
            preview_data['sections'].append({
                'title': 'Beneficiary Details',
                'type': 'table',
                'columns': [
                    {'id': 'name', 'name': 'Name'},
                    {'id': 'email', 'name': 'Email'},
                    {'id': 'status', 'name': 'Status'},
                    {'id': 'test_score', 'name': 'Avg Score'},
                    {'id': 'trainer', 'name': 'Trainer'}
                ],
                'data': [
                    {
                        'name': 'John Doe',
                        'email': 'john.doe@example.com',
                        'status': 'Active',
                        'test_score': '85%',
                        'trainer': 'Sarah Johnson'
                    },
                    {
                        'name': 'Jane Smith',
                        'email': 'jane.smith@example.com',
                        'status': 'Active',
                        'test_score': '92%',
                        'trainer': 'Mike Wilson'
                    }
                ]
            })
            
        elif data['type'] == 'program':
            # Chart section
            preview_data['sections'].append({
                'title': 'Program Enrollment Trends',
                'type': 'chart',
                'chartType': 'line',
                'data': {}  # Chart data would go here
            })
            
            # Table section
            preview_data['sections'].append({
                'title': 'Program Details',
                'type': 'table',
                'columns': [
                    {'id': 'name', 'name': 'Program'},
                    {'id': 'enrollment', 'name': 'Enrollments'},
                    {'id': 'completion', 'name': 'Completion Rate'},
                    {'id': 'attendance', 'name': 'Attendance'}
                ],
                'data': [
                    {
                        'name': 'Leadership Training',
                        'enrollment': '25',
                        'completion': '88%',
                        'attendance': '94%'
                    },
                    {
                        'name': 'Technical Skills',
                        'enrollment': '40',
                        'completion': '76%',
                        'attendance': '89%'
                    }
                ]
            })
        
        return jsonify(preview_data), 200
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@reports_bp.route('/reports/generate', methods=['POST'])
@jwt_required()
def generate_report():
    """Generate a report in the requested format."""
    user_id = get_jwt_identity()
    user = User.query.get(user_id)
    data = request.get_json()
    
    try:
        # Create temporary report
        report = Report(
            name=data.get('name', 'Generated Report'),
            description=data.get('description', ''),
            type=data['type'],
            format=data.get('format', 'pdf'),
            parameters=data,
            created_by_id=user_id,
            tenant_id=user.tenant_id,
            status='generating'
        )
        
        db.session.add(report)
        db.session.commit()
        
        # Generate report data
        if report.type == 'beneficiary':
            report_data = generate_beneficiary_report(data)
        elif report.type == 'trainer':
            report_data = generate_trainer_report(data)
        elif report.type == 'program':
            report_data = generate_program_report(data)
        elif report.type == 'performance':
            report_data = generate_performance_report(data)
        else:
            report_data = generate_general_report(data)
        
        # Generate file based on format
        if data.get('format') == 'pdf':
            file_path = generate_pdf_report(report_data, report)
        elif data.get('format') == 'xlsx':
            file_path = generate_excel_report(report_data, report)
        else:  # csv
            file_path = generate_csv_report(report_data, report)
        
        # Update report
        report.status = 'completed'
        report.file_path = file_path
        report.file_size = os.path.getsize(file_path)
        report.last_generated = datetime.now()
        db.session.commit()
        
        # Return file
        return send_file(
            file_path,
            as_attachment=True,
            download_name=f"{report.name}_{datetime.now().strftime('%Y%m%d')}.{report.format}"
        )
        
    except Exception as e:
        if 'report' in locals():
            report.status = 'failed'
            db.session.commit()
        return jsonify({'error': str(e)}), 500